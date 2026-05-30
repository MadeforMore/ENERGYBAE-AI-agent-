from __future__ import annotations

import os
import shutil
import tempfile
from datetime import datetime
from pathlib import Path
from statistics import mean
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.config import settings
from app.models.schema import BillExtraction
from app.services.template_mapping import PANEL_WATTAGE_CELL, TEMPLATE_SLOTS, TemplateSlot


SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def build_workbook(bills: list[BillExtraction], panel_wattage: float) -> Path:
    return build_workbooks(bills, panel_wattage)


def build_workbooks(bills: list[BillExtraction], panel_wattage: float) -> Path:
    workbooks: list[Path] = []

    for start in range(0, len(bills), len(TEMPLATE_SLOTS)):
        chunk = bills[start : start + len(TEMPLATE_SLOTS)]
        sequence = len(workbooks) + 1
        workbooks.append(build_single_workbook(chunk, panel_wattage, sequence))

    if len(workbooks) == 1:
        return workbooks[0]

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    zip_path = settings.output_dir / f"energybae_batch_{timestamp}.zip"
    with ZipFile(zip_path, "w", compression=ZIP_DEFLATED) as archive:
        for workbook_path in workbooks:
            archive.write(workbook_path, workbook_path.name)
    return zip_path


def build_single_workbook(bills: list[BillExtraction], panel_wattage: float, sequence: int) -> Path:
    workbook = load_workbook(settings.template_path)
    sheet = workbook[workbook.sheetnames[0]]
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    sheet[PANEL_WATTAGE_CELL] = panel_wattage

    cached_values: dict[str, float | None] = {}
    slot_results: list[dict[str, float | None]] = []

    for index, slot in enumerate(TEMPLATE_SLOTS):
        bill = bills[index] if index < len(bills) else None
        clear_slot(sheet, slot)
        if bill is not None:
            fill_slot(sheet, slot, bill)
            slot_result = calculate_slot_values(slot, bill, panel_wattage)
            slot_results.append(slot_result)
            cached_values.update(slot_result)
        else:
            slot_results.append({})

    cached_values.update(calculate_total_values(slot_results))

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    suffix = f"_{sequence}" if len(bills) > 0 and sequence > 0 else ""
    output_path = settings.output_dir / f"energybae_output_{timestamp}{suffix}.xlsx"
    workbook.save(output_path)
    patch_cached_formula_values(output_path, cached_values)
    return output_path


def clear_slot(sheet: Worksheet, slot: TemplateSlot) -> None:
    for cell_ref in [
        slot.consumer_name,
        slot.consumer_number,
        slot.fixed_charges,
        slot.sanctioned_load,
        slot.connection_type,
        slot.bill_amount,
    ]:
        sheet[cell_ref] = None

    for cell_ref in slot.month_labels + slot.units:
        sheet[cell_ref] = None


def fill_slot(sheet: Worksheet, slot: TemplateSlot, bill: BillExtraction) -> None:
    sheet[slot.consumer_name] = bill.consumer_name
    sheet[slot.consumer_number] = str(bill.consumer_number or "")
    sheet[slot.fixed_charges] = bill.fixed_charges
    sheet[slot.sanctioned_load] = format_kw(bill)
    sheet[slot.connection_type] = bill.connection_type or bill.tariff_category
    sheet[slot.bill_amount] = bill.excel_bill_amount

    for idx, history in enumerate(bill.monthly_history[:12]):
        month_cell = slot.month_labels[idx]
        units_cell = slot.units[idx]
        sheet[month_cell] = month_label_to_excel_date(history.month_label, history.month_iso)
        sheet[units_cell] = history.units


def month_label_to_excel_date(month_label: str, month_iso: str | None):
    if month_iso:
        return datetime.strptime(month_iso + "-01", "%Y-%m-%d")
    if month_label:
        return datetime.strptime(month_label, "%b %Y")
    return None


def format_kw(bill: BillExtraction) -> str | float | None:
    if bill.sanctioned_load_text:
        return bill.sanctioned_load_text
    if bill.sanctioned_load_kw is None:
        return None
    return f"{bill.sanctioned_load_kw:.2f} KW"


def calculate_slot_values(slot: TemplateSlot, bill: BillExtraction, panel_wattage: float) -> dict[str, float | None]:
    units = [item.units for item in bill.monthly_history[:12] if item.units is not None]
    average_units = mean(units) if units else None
    current_units = bill.units_consumed if bill.units_consumed is not None else (units[-1] if units else None)
    fixed_charges = bill.fixed_charges
    bill_amount = bill.excel_bill_amount

    if current_units in (None, 0) or bill_amount is None or fixed_charges is None:
        unit_cost = None
    else:
        unit_cost = (bill_amount - fixed_charges) / current_units

    kw = (average_units * 12 * 1.1) / 1400 if average_units is not None else None
    solar_panels = (kw / panel_wattage * 1000) if kw is not None and panel_wattage else None
    solar_capacity = (round(solar_panels, 0) * panel_wattage / 1000) if solar_panels is not None else None
    number_of_panels = (solar_capacity / panel_wattage * 1000) if solar_capacity is not None and panel_wattage else None

    if slot.name == "left":
        return {
            "D22": average_units,
            "E22": bill_amount,
            "F22": unit_cost,
            "D23": kw,
            "D24": solar_panels,
            "D25": solar_capacity,
            "D26": number_of_panels,
            "F20": unit_cost,
        }

    return {
        "H22": average_units,
        "I22": bill_amount,
        "J22": unit_cost,
        "H23": kw,
        "H24": solar_panels,
        "H25": solar_capacity,
        "H26": number_of_panels,
        "J20": unit_cost,
    }


def calculate_total_values(slot_results: list[dict[str, float | None]]) -> dict[str, float | None]:
    capacities = [result.get("D25") for result in slot_results if result.get("D25") is not None]
    capacities.extend(result.get("H25") for result in slot_results if result.get("H25") is not None)
    panels = [result.get("D26") for result in slot_results if result.get("D26") is not None]
    panels.extend(result.get("H26") for result in slot_results if result.get("H26") is not None)

    return {
        "D29": sum(capacities) if capacities else None,
        "D30": sum(panels) if panels else None,
    }


def patch_cached_formula_values(workbook_path: Path, cached_values: dict[str, float | None]) -> None:
    temp_fd, temp_name = tempfile.mkstemp(suffix=".xlsx")
    os.close(temp_fd)
    Path(temp_name).unlink(missing_ok=True)
    temp_path = Path(temp_name)

    with ZipFile(workbook_path, "r") as source, ZipFile(temp_path, "w", compression=ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = update_sheet_formula_cache(data, cached_values)
            elif item.filename == "xl/workbook.xml":
                data = ensure_recalc_flags(data)
            target.writestr(item, data)

    shutil.move(temp_path, workbook_path)


def ensure_recalc_flags(xml_bytes: bytes) -> bytes:
    ns = {"a": SPREADSHEET_NS}
    ET.register_namespace("", SPREADSHEET_NS)
    root = ET.fromstring(xml_bytes)
    calc = root.find("a:calcPr", ns)
    if calc is None:
        calc = ET.SubElement(root, f"{{{SPREADSHEET_NS}}}calcPr")
    calc.set("calcMode", "auto")
    calc.set("fullCalcOnLoad", "1")
    calc.set("forceFullCalc", "1")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def update_sheet_formula_cache(xml_bytes: bytes, cached_values: dict[str, float | None]) -> bytes:
    ns = {"a": SPREADSHEET_NS}
    ET.register_namespace("", SPREADSHEET_NS)
    root = ET.fromstring(xml_bytes)

    for cell in root.findall(".//a:c", ns):
        ref = cell.attrib.get("r")
        if ref not in cached_values:
            continue
        formula = cell.find("a:f", ns)
        if formula is None:
            continue
        value_node = cell.find("a:v", ns)
        if value_node is None:
            value_node = ET.SubElement(cell, f"{{{SPREADSHEET_NS}}}v")
        value = cached_values[ref]
        value_node.text = "" if value is None else format_number(value)

    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def format_number(value: float) -> str:
    return f"{value:.10f}".rstrip("0").rstrip(".")
